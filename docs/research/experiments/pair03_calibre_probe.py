#!/usr/bin/env python3
"""R11 synthetic Calibre DOCX -> EPUB -> DOCX fidelity probe.

Requires installed free Calibre ebook-convert + python-docx. It measures, but
does not guarantee, semantic formatting and author-quality import fidelity.
Test documents are synthetic and discarded after the local CI job.
"""
import json
import re
import subprocess
import tempfile
from pathlib import Path
from docx import Document

TEXTS = [
    "Mara held the lantern beside the river.",
    "Elias whispered the secret at dawn.",
    "The courier returned before sunset.",
]
TITLES = ["Chapter One", "Chapter Two", "Chapter Three"]


def convert(src, dst):
    p = subprocess.run(["ebook-convert", str(src), str(dst)], capture_output=True,
                       text=True, timeout=180, check=False)
    if p.returncode:
        raise RuntimeError("Calibre conversion failed: " + p.stderr[-800:])
    return p.stderr[-1200:]


def run():
    version = subprocess.run(["ebook-convert", "--version"], capture_output=True,
                             text=True, timeout=20, check=True).stdout.strip()
    with tempfile.TemporaryDirectory(prefix="vellistry-r11-synthetic-") as folder:
        root = Path(folder)
        src = root / "fiction.docx"
        epub = root / "stage.epub"
        target = root / "returned.docx"
        source = Document()
        for title, text in zip(TITLES, TEXTS):
            source.add_heading(title, level=1)
            para = source.add_paragraph()
            para.add_run(text.split(" ")[0] + " ").bold = True
            para.add_run(" ".join(text.split(" ")[1:])).italic = True
        source.save(src)
        warning_forward = convert(src, epub)
        warning_return = convert(epub, target)
        if not epub.exists() or not target.exists():
            raise AssertionError("Calibre failed to emit output")
        returned = Document(target)
        parts = [p.text for p in returned.paragraphs]
        whole = "\n".join(parts)
        observed = {}
        for title, text in zip(TITLES, TEXTS):
            title_ok = title in whole
            # Whitespace may change between the two formats; textual content must not.
            sentence_ok = re.sub(r"\s+", " ", text) in re.sub(r"\s+", " ", whole)
            observed[title] = {"heading_text_found": title_ok, "sentence_found": sentence_ok}
            if not title_ok or not sentence_ok:
                raise AssertionError("DOCX roundtrip changed a source heading/sentence: " + title)
        original_headings = len([p for p in source.paragraphs if p.style.name.startswith("Heading")])
        final_headings = len([p for p in returned.paragraphs if p.style.name.startswith("Heading")])
        result = {
            "fixture": "Synthetic three-chapter DOCX, bold + italic source spans",
            "ebook_convert_version": version,
            "original_heading_count": original_headings,
            "returned_heading_count": final_headings,
            "heading_style_count_equal": original_headings == final_headings,
            "all_source_headings_and_sentences_found": True,
            "per_chapter": observed,
            "output_file_byte_sizes": {"epub": epub.stat().st_size, "docx": target.stat().st_size},
            "warnings_detected": {"forward": bool(warning_forward.strip()),
                                  "return": bool(warning_return.strip())},
            "unverified": [
                "Exact paragraph/run formatting preservation and independent Microsoft Word rendering",
                "Page breaks, scene breaks, footnotes, tracked changes and comments",
                "Unicode beyond controlled fixture, hostile OOXML sandboxing",
                "Import conversion into a Tiptap 3 editor and a full author-project export",
            ],
        }
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    run()
