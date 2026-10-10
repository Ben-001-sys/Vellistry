"""Option B Wave 2 R11: semantic subset Tiptap JSON <-> Word DOCX research.

Uses free python-docx; does NOT replace shipping importer/exporter or prove
full Microsoft Word fidelity. Unsupported features FAIL CLOSED without silent
drops. No user-provided manuscripts or network operations.
"""
from __future__ import annotations
from copy import deepcopy
from io import BytesIO
from docx import Document
from docx.oxml.ns import qn

ALLOWED_MARKS={"bold","italic","strike"}
MAX_BLOCKS=120
MAX_CHARS=120000

class UnsupportedFormat(ValueError):
    pass

def validate(doc):
    if not isinstance(doc,dict) or set(doc)!={"type","content"} or doc["type"]!="doc":
        raise UnsupportedFormat("document envelope unsupported")
    blocks=doc["content"]
    if not isinstance(blocks,list) or len(blocks)>MAX_BLOCKS:
        raise UnsupportedFormat("too many blocks")
    size=0
    for b in blocks:
        if not isinstance(b,dict) or b.get("type") not in ("heading","paragraph"):
            raise UnsupportedFormat("unsupported block")
        typ=b["type"]
        if typ=="heading":
            if set(b)!={"type","attrs","content"} or b["attrs"] not in ({"level":1},{"level":2},{"level":3}):
                raise UnsupportedFormat("unsupported heading")
        elif set(b) not in ({"type","content"},{"type"}):
            raise UnsupportedFormat("unsupported paragraph attributes")
        runs=b.get("content",[])
        if not isinstance(runs,list):raise UnsupportedFormat("invalid runs")
        for t in runs:
            if not isinstance(t,dict) or t.get("type")!="text" or set(t)-{"type","text","marks"}:
                raise UnsupportedFormat("unsupported inline node")
            text=t.get("text")
            if not isinstance(text,str) or not text:
                raise UnsupportedFormat("invalid run")
            size+=len(text)
            marks=t.get("marks",[])
            if not isinstance(marks,list) or len({m.get("type") for m in marks if isinstance(m,dict)})!=len(marks):
                raise UnsupportedFormat("invalid/duplicate marks")
            if any(not isinstance(m,dict) or set(m)!={"type"} or m["type"] not in ALLOWED_MARKS for m in marks):
                raise UnsupportedFormat("unsupported marks/attrs")
    if size>MAX_CHARS:raise UnsupportedFormat("document exceeds research size limit")
    return deepcopy(doc)

def docx_export(document):
    src=validate(document)
    out=Document()
    for block in src["content"]:
        typ=block["type"]
        paragraph=out.add_heading(level=block["attrs"]["level"]) if typ=="heading" else out.add_paragraph()
        for token in block.get("content",[]):
            run=paragraph.add_run(token["text"])
            kinds={m["type"] for m in token.get("marks",[])}
            run.bold="bold" in kinds
            run.italic="italic" in kinds
            run.font.strike="strike" in kinds
    data=BytesIO()
    out.save(data)
    return data.getvalue()

# A paragraph's .runs omits nested hyperlinks, field results, revision wrappers,
# and other Word XML. Reject unrepresented structures BEFORE extracting any text.
# This is deliberately a narrow research subset, not a general OOXML sanitizer.
_SUPPORTED_RUN_FORMATS={qn("w:b"),qn("w:i"),qn("w:strike")}
_SUPPORTED_BOOLEAN_VALUES={"0","1","true","false","on","off"}

def _validate_word_structure(source):
    for element in source.element.body.iterchildren():
        if element.tag not in (qn("w:p"),qn("w:sectPr")):
            raise UnsupportedFormat("unsupported Word body element")
        if element.tag==qn("w:sectPr"):
            if any(child.tag in (qn("w:headerReference"),qn("w:footerReference"))
                   for child in element.iterchildren()):
                raise UnsupportedFormat("Word headers/footers unsupported")
            continue
        for child in element.iterchildren():
            if child.tag==qn("w:pPr"):
                if any(prop.tag!=qn("w:pStyle") for prop in child.iterchildren()):
                    raise UnsupportedFormat("unsupported Word paragraph properties")
            elif child.tag==qn("w:r"):
                for part in child.iterchildren():
                    if part.tag==qn("w:rPr"):
                        for prop in part.iterchildren():
                            if (prop.tag not in _SUPPORTED_RUN_FORMATS
                                or set(prop.attrib)-{qn("w:val")}
                                or prop.get(qn("w:val"),"1") not in _SUPPORTED_BOOLEAN_VALUES):
                                raise UnsupportedFormat("unsupported Word run formatting")
                    elif part.tag!=qn("w:t"):
                        raise UnsupportedFormat("unsupported Word run content")
            else:
                raise UnsupportedFormat("unsupported nested Word paragraph content")

def docx_import(data):
    if not isinstance(data,bytes) or len(data)>2_000_000:
        raise UnsupportedFormat("invalid/oversize DOCX")
    try:
        source=Document(BytesIO(data))
    except Exception as exc:
        raise UnsupportedFormat("invalid DOCX package") from exc
    _validate_word_structure(source)
    blocks=[]
    for p in source.paragraphs:
        style=p.style.name
        if style in ("Normal","Body Text"):
            typ="paragraph"; attrs=None
        elif style in ("Heading 1","Heading 2","Heading 3"):
            typ="heading";attrs={"level":int(style.split()[-1])}
        else:
            raise UnsupportedFormat("unsupported Word paragraph style "+style)
        tokens=[]
        for run in p.runs:
            if not run.text:
                continue
            if run.font.strike is True:
                struck=True
            else:
                struck=False
            marks=[{"type":name} for name,enabled in (
                ("bold",run.bold is True),
                ("italic",run.italic is True),
                ("strike",struck)) if enabled]
            tokens.append({"type":"text","text":run.text,**({"marks":marks} if marks else {})})
        block={"type":typ, **({"attrs":attrs} if attrs else {})}
        if tokens:block["content"]=tokens
        elif typ=="heading":block["content"]=[]
        blocks.append(block)
    if len(source.tables)>0 or len(source.inline_shapes)>0:
        raise UnsupportedFormat("tables/images require explicit supported import")
    return validate({"type":"doc","content":blocks})

def semantic(doc):
    """Compare blocks/marks independent of implicit False stored as None."""
    r=validate(doc)
    for b in r["content"]:
        if not b.get("content"):
            b.pop("content",None)
        for t in b.get("content",[]):
            if t.get("marks"):
                t["marks"]=sorted(t["marks"],key=lambda v:v["type"])
            else:t.pop("marks",None)
    return r

def fixture():
    return {"type":"doc","content":[
        {"type":"heading","attrs":{"level":1},"content":[{"type":"text","text":"Chapter One"}]},
        {"type":"paragraph","content":[
            {"type":"text","text":"Mara saw "},
            {"type":"text","text":"a blue lantern", "marks":[{"type":"bold"}]},
            {"type":"text","text":" and whispered "},
            {"type":"text","text":"stay", "marks":[{"type":"italic"}]},
            {"type":"text","text":". 雪 في السماء. “Return!”"}]},
        {"type":"heading","attrs":{"level":2},"content":[{"type":"text","text":"Scene Two"}]},
        {"type":"paragraph","content":[{"type":"text","text":"She crossed out "},
            {"type":"text","text":"perhaps", "marks":[{"type":"strike"}]},
            {"type":"text","text":" and continued."}]},
        {"type":"paragraph"}
    ]}
