"""R11 Wave2 free DOCX fidelity: semantic heading/mark preservation.
No file is used outside a disposable in-memory synthetic fixture.
"""
import unittest
from copy import deepcopy
from io import BytesIO
from docx import Document
from pair_wave02_docx_semantics import (
    docx_export, docx_import, semantic, validate, UnsupportedFormat, fixture
)
from pair03_portability import backup_bytes, restore_bytes, synthetic_project, UnsafeArchive

class DocxRoundtripTests(unittest.TestCase):
    def test_01_heading_one_and_two_word_style_survive(self):
        original=fixture(); data=docx_export(original); word=Document(BytesIO(data))
        headings=[p.style.name for p in word.paragraphs if p.style.name.startswith("Heading")]
        self.assertEqual(headings,["Heading 1","Heading 2"])
        self.assertEqual(semantic(docx_import(data)),semantic(original))

    def test_02_bold_italic_strikethrough_run_semantics(self):
        expected=semantic(fixture())
        actual=semantic(docx_import(docx_export(fixture())))
        self.assertEqual(actual["content"][1]["content"][1]["marks"],[{"type":"bold"}])
        self.assertEqual(actual["content"][1]["content"][3]["marks"],[{"type":"italic"}])
        self.assertEqual(actual["content"][3]["content"][1]["marks"],[{"type":"strike"}])
        self.assertEqual(actual,expected)

    def test_03_multilingual_and_curly_quotes_survive(self):
        r=docx_import(docx_export(fixture()))
        full=" ".join(t["text"] for p in r["content"] for t in p.get("content",[]))
        for snippet in ("雪","في السماء","“Return!”","Mara"):
            self.assertIn(snippet,full)

    def test_04_empty_paragraph_stays_empty(self):
        restored=docx_import(docx_export(fixture()))
        self.assertEqual(restored["content"][-1],{"type":"paragraph"})

    def test_05_input_not_mutated(self):
        x=fixture();before=deepcopy(x)
        docx_import(docx_export(x))
        self.assertEqual(x,before)

    def test_06_reject_unsupported_content_without_partial_export(self):
        x=fixture();x["content"].append({"type":"image","attrs":{"src":"https://elsewhere.example"}})
        with self.assertRaises(UnsupportedFormat):
            docx_export(x)

    def test_07_reject_unknown_html_like_marks(self):
        x=fixture();x["content"][1]["content"][1]["marks"]=[{"type":"link","attrs":{"href":"https://x"}}]
        with self.assertRaises(UnsupportedFormat):
            docx_export(x)

    def test_08_reject_foreign_docx_table_instead_of_dropping_it(self):
        d=Document();d.add_paragraph("Text");d.add_table(rows=1,cols=1)
        out=BytesIO();d.save(out)
        with self.assertRaises(UnsupportedFormat):
            docx_import(out.getvalue())

    def test_09_bad_docx_rejected(self):
        with self.assertRaises(UnsupportedFormat):
            docx_import(b"not a ZIP at all")

    def test_10_two_book_archive_restores_original_full_rich_document(self):
        p=synthetic_project()
        p["books"][0]["chapters"][0]["document"]=fixture()
        p["books"][0]["chapters"][0]["revision"]=17
        existing={"untouched":True}
        restored=restore_bytes(backup_bytes(p),existing)
        self.assertEqual(restored,p)
        self.assertEqual(restored["books"][0]["chapters"][0]["revision"],17)
        self.assertEqual(restored["books"][0]["profiles"][1]["status"],"pending_not_canon")
        self.assertEqual(restored["books"][0]["rules"][0]["enabled"],False)
        self.assertEqual(existing,{"untouched":True})

    def test_11_archive_backup_keeps_separate_author_values(self):
        restored=restore_bytes(backup_bytes(synthetic_project()),{})
        self.assertEqual(restored["books"][0]["owner"],"alice")
        self.assertEqual(restored["books"][1]["owner"],"bob")

    def test_12_unsupported_paragraph_style_docx_fails_closed(self):
        doc=Document(); doc.add_paragraph("Special",style="Title")
        output=BytesIO(); doc.save(output)
        with self.assertRaises(UnsupportedFormat):
            docx_import(output.getvalue())

if __name__=="__main__":unittest.main()
